from openpyxl import load_workbook
from xpaths import *

excel = load_workbook(r"C:\Users\User\Documents\Ameya Data.xlsx")
sheet = excel['Sheet2']

NewUser = AccountCreation()
NewUser.onboarding()
NewUser.userlogin()

for i in range(2, sheet.max_row+1):
    NewUser.memberdropdown()

    NewUser.nametab(title=sheet.cell(row=i, column=1).value, fname=sheet.cell(row=i, column=2).value, sname=sheet.cell(row=i, column=3).value, dob=sheet.cell(row=i, column=4).value, gender=sheet.cell(row=i, column=5).value, maiden=sheet.cell(row=i, column=6).value, nick=sheet.cell(row=i, column=7).value, ptype=sheet.cell(row=i, column=8).value, idnumber=sheet.cell(row=i, column=9).value, isdate=sheet.cell(row=i, column=10).value, exdate=sheet.cell(row=i, column=11).value, country=sheet.cell(row=i, column=12).value)
    NewUser.communicationtab(country=sheet.cell(row=i, column=13).value, state=sheet.cell(row=i, column=14).value, city=sheet.cell(row=i, column=15).value, zipcode=sheet.cell(row=i, column=16).value)
    NewUser.controltab(person=sheet.cell(row=i, column=17).value)

    memberidnum = NewUser.memberid()
    sheet.cell(row=i, column=19).value = str(memberidnum)
    print(f'member id is {memberidnum}')
    excel.save(r"C:\Users\User\Documents\Ameya Data.xlsx")

    print(f'done for {i}th row, thank you!!!!!')
    print('*' * 60)
